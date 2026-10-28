import openpyxl, json, sys
sys.path.insert(0,'.')
INV={'Atacar':'atk','Mover':'mov','Curar':'heal','Magia':'magic'}
def read_cards(path):
    wb=openpyxl.load_workbook(path,data_only=True)
    heroes={}; monsters={}
    for r in list(wb['Cartas Heróis'].iter_rows(values_only=True))[1:]:
        if not r[0]: continue
        c={'n':str(r[3]),'t':[INV[r[4]],int(r[5])],'b':[INV[r[6]],int(r[7])]}
        if len(r)>8 and str(r[8] or '').strip().lower().startswith('s'): c['k']=1
        heroes.setdefault(r[0],[]).append(c)
    for r in list(wb['Cartas Monstros'].iter_rows(values_only=True))[1:]:
        if not r[0]: continue
        c={'n':str(r[2]),'t':[INV[r[3]],int(r[4])],'b':[INV[r[5]],int(r[6])]}
        if str(r[7]).strip().lower().startswith('s'): c['s']=1
        monsters.setdefault(r[0],[]).append(c)
    for k,v in monsters.items(): assert any(c.get('s') for c in v),('sem carta de embaralhamento',k)
    rg={}
    if 'Alcance Monstros' in wb.sheetnames:
        for r in list(wb['Alcance Monstros'].iter_rows(values_only=True))[1:]:
            if r[0] and r[1] is not None: rg[str(r[0])]=int(r[1])
    return {'heroes':heroes,'monsters':monsters,'range':rg}
def build(xlsx='cartas.xlsx',tpl='tpl.html',out='catacumba.html'):
    cards=read_cards(xlsx)
    t=open(tpl).read()
    assert '/*__MONSTERS__*/[]' in t and '/*__CARDS__*/{}' in t
    t=t.replace('/*__MONSTERS__*/[]',open('monsters.json').read()).replace('/*__CARDS__*/{}',json.dumps(cards,ensure_ascii=False,separators=(',',':')))
    open(out,'w').write(t); return cards
if __name__=='__main__':
    c=build(); print(len(c['heroes']),'heróis',len(c['monsters']),'monstros')
