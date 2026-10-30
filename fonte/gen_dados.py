import openpyxl, json, re
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
A='/mnt/user-data/attachments/02b41faa-66e7-503e-a967-caf8268844b7/'
wb=openpyxl.load_workbook(A+'9864fbc8-cartas.xlsx')
arm=openpyxl.load_workbook(A+'279c6876-Armas.xlsx',data_only=True).active
mons={m['n']:m for m in json.load(open('monsters.json'))}
HD=PatternFill('solid',fgColor='2D323C'); HF=Font(bold=True,color='FFFFFF')
def head(w):
    for c in range(1,w.max_column+1):
        w.cell(1,c).fill=HD; w.cell(1,c).font=HF
    w.freeze_panes='A2'
def widths(w,ws):
    for i,x in enumerate(ws): w.column_dimensions[get_column_letter(i+1)].width=x
ACT={'A':'Atacar','M':'Mover','C':'Curar','G':'Magia','S':'Saquear'}
FXN={'perf':'Perfurar','area':'Área','emp':'Empurrar','ferir':'Ferir','atord':'Atordoar','gelo':'Congelar','esc':'Escudo','ret':'Retaliar','carea':'Cura em área','salto':'Salto','dren':'Drenar'}
def fxs(t):
    o=[]
    for e in (t.split(':')[1].split(',') if ':' in t else []):
        m=re.match(r'([a-z]+)(\d*)$',e); k,n=m.group(1),m.group(2); assert k in FXN,e
        o.append(FXN[k]+(' '+n if n else ''))
    return '; '.join(o)
def card(s):
    p=s.split('|'); a,b=p[1].split(':')[0],p[2].split(':')[0]
    return (p[0],ACT[a[0]],int(a[1:]),ACT[b[0]],int(b[1:]),'Sim' if len(p)>3 else 'Não',fxs(p[1]),fxs(p[2]))

# ---------------- decks por classe (7 cartas) ----------------
DECKS={
'Espadachim':['Golpe Firme|A3|M1','Avanço|A3|M1:salto','Corte e Saque|A2|S1','Guarda Aberta|A2|M0:esc1','Estocada|A3:perf1|A1','Reposicionar|A2|M1','Investida|A4|M0'],
'Guerreiro':['Machadada Saqueadora|A3|S1','Muralha de Escudos|A2|M1:esc2','Pancada Giratória|A2:area|A2','Investida de Escudo|A2:emp1|M1:esc1','Golpe Brutal|A4|M0','Postura de Retaliação|A1|M1:esc1,ret2','Avançar e Bater|A3:perf1|M1'],
'Samurai':['Iaijutsu|A4:perf1|M0','Passo Silencioso|A2|M2','Corte Cruzado|A3:area|A1','Postura Alta|A2|M1:ret1','Lâmina Veloz|A3|S1','Contra-Golpe|A3:perf2|M1','Fio da Navalha|A2:perf1|A2:ferir'],
'Ninja':['Golpe Sombrio|A2|S1','Salto|A1|M1:salto','Kunai Envenenada|A2:ferir|M1','Mãos Leves|A2:perf1|S1','Corte Duplo|A2|A2','Passo das Sombras|A2|M2:salto|B','Furto Letal|A3:perf2|S1'],
'Cavaleiro':['Carga|A3:emp1|M1','Lança Firme|A3|M0:esc1','Carga Total|A3|M2:salto|B','Estocada Montada|A2:perf1|M1','Reposicionar|A2|M0','Investida e Saque|A3|S1','Guarda|A1|M1:esc2'],
'Cavaleiro Alado':['Mergulho Rapineiro|A2|S1','Voo Rasante|A2|M1','Ataque Aéreo|A3|M0','Voo Supremo|A2|M2|B','Lança do Céu|A3:perf1|M1','Giro no Ar|A3:area|M0','Pouso Certeiro|A2|A1:esc1'],
'Cavaleiro a Vapor':['Pancada de Aço|A3:emp1|M0','Vapor Quente|A2:ferir|M1','Avanço Blindado|A2|M1:esc2','Pistão|A4|S1','Marcha Lenta|A1:esc1|M1','Esmagar|A3:area|A1','Escudo de Vapor|A1|M1:esc3'],
'Cavaleiro de Assalto':['Tiro Pesado|A3:perf1|M0','Reposicionar|A2|M1','Barragem|A2:area|A2','Mira Firme|A3|M1','Recarga|A1:esc1|M1','Disparo e Saque|A2|S1','Bombardeio|A4:area|M0'],
'Arqueiro':['Flechada|A2|M1','Recuar e Atirar|A2|M1','Tiro Certeiro|A3:perf1|M0','Chuva de Flechas|A2:area|A2','Posicionar|A1|M2','Disparo Veloz|A1|S1','Flecha Envenenada|A3:ferir|M1'],
'Mago':['Bola de Fogo|G3:area|M0','Passo Arcano|G2|M1:salto','Raio Frio|G2:gelo|M1','Barreira Arcana|A1|M1:esc2','Explosão|G4:area|M0|B','Fluxo de Mana|G2|G1','Cajado Cobiçoso|A2|S1'],
'Sacerdote':['Prece|C2|M1','Apoio em Combate|A1|C2:esc1','Bênção Coletiva|C3:carea|M0','Luz Restauradora|A1|C2:esc1','Passo Sagrado|M1|C1','Cajado Sagrado|A2|S1','Cura Maior|C4|M0|B'],
'Monge':['Soco Ágil|A2|S1','Palma Aberta|A3:emp1|M0','Respiração|C1|M1:esc1','Passo Leve|A1|M2','Combo|A2|A2','Chute Giratório|A3:area|M1','Meditação|A1:dren1|C2'],
'Aviário':['Rasante|A2|S1','Bicada|A2:perf1|M1','Revoada|A2|M2|B','Garras|A3|M0','Mergulho|A3|M1','Asas Abertas|A2:emp1|M0','Pouso e Coleta|A1|S1'],
'Lobisomen':['Mordida|A3:dren1|M1','Salto Selvagem|A2|M1:salto','Garras Duplas|A2:ferir|A2','Uivo|A1:atord|M1','Investida Feroz|A4|M0','Caçada Selvagem|A2|M2:salto|B','Dilacerar|A3:ferir|S1'],
'Dragão':['Sopro|G3:area|M0','Voo|A1|M1','Garra Avarenta|A2|S1','Rabada|A2:emp1|A1','Mergulho|A2|M1','Rugido|A1:atord|M1','Baforada|G2:ferir|M1'],
'Robô':['Soco de Aço|A3|S1','Marcha|A1:esc1|M1','Canhão de Braço|A2:perf1|A2','Blindagem|A1|M1:esc3','Esmagar|A4:emp1|M0','Avanço Pesado|A2|M1:esc1','Pulso|A2:area|M1'],
'Mascote':['Tropeção|A1:emp1|M1','Fuga Cômica|M1|M1:esc1','Cutucada|A2|S1','Pulinho|A1|M1:salto','Sorte Grande|A3|M0','Dança Gananciosa|A1|S1','Susto|A1:atord|A1'],
'Criatura Mágica':['Faísca|G2|M1','Flutuar|G1|M1','Pulso Mágico|G3:area|M0','Névoa|A1|M1:esc1','Raio|G2:gelo|G1','Deslize|A1|M1:salto','Carga Mística|G3|S1'],
}
SIG={
'max':'Lâmina do Líder|A4:perf1|M1:esc1','tao':'Blaze|G5:area|M1|B','anri':'Freeze|G4:gelo|M2|B','alef':'Bolt|G5:emp1|M0|B',
'lowe':'Mãos Curadoras|C3:esc1|M2','khris':'Luz de Kyantol|C4:carea|M1','torasu':'Milagre|C5:carea|M0|B',
'luke':'Muralha Viva|A3|A1:esc2','gort':'Machado Furioso|A4:area|M0',
'ken':'Galope Veloz|A2|M2:salto','mae':'Estocada Veloz|A3:perf1|M2','arthur':'Lança Sagrada|A3|C2:esc1','pelle':'Carga Equilibrada|A3|M1:esc1','vankar':'Escudo Encouraçado|A3|M0:esc3,ret1','earnest':'Golpe Esmagador|A4:emp1|A1',
'hans':'Tiro Perfurante|A4:perf2|M0','diane':'Flecha Gélida|A3:gelo|M2','gong':'Punho Curador|A3:dren2|C2',
'amon':'Voo Ágil|A2:perf1|M2','balbaroy':'Garra do Céu|A4|M1:emp1','zylo':'Uivo Selvagem|A4:atord|M2','kokichi':'Voo Relâmpago|A2|M2',
'guntz':'Fortaleza a Vapor|A4|M0:esc3','musashi':'Corte da Lua|A4:area|M1','hanzou':'Lâmina Sombria|A3:ferir|M2','lyle':'Canhão Pesado|A5:area|M0',
'bleu':'Sopro do Dragão|G4:area|M1|B','adam':'Soco Atômico|A5:emp2|M0|B','jogurt':'Sorte do Iogurte|A4|M2','domingo':'Rajada Mística|G4:gelo|M1|B',
}
BASE={'max':'Espadachim','tao':'Mago','anri':'Mago','alef':'Mago','lowe':'Sacerdote','khris':'Sacerdote','torasu':'Sacerdote','luke':'Guerreiro','gort':'Guerreiro',
'ken':'Cavaleiro','mae':'Cavaleiro','arthur':'Cavaleiro','pelle':'Cavaleiro','vankar':'Cavaleiro','earnest':'Cavaleiro','hans':'Arqueiro','diane':'Arqueiro','gong':'Monge',
'amon':'Aviário','balbaroy':'Aviário','zylo':'Lobisomen','kokichi':'Cavaleiro Alado','domingo':'Criatura Mágica','guntz':'Cavaleiro a Vapor','musashi':'Samurai','hanzou':'Ninja',
'lyle':'Cavaleiro de Assalto','bleu':'Dragão','adam':'Robô','jogurt':'Mascote'}
# PV, Movimento, ATQ base, DF, AG, Voa
ST={'max':(12,2,1,0,8,0),'tao':(9,1,1,0,6,0),'anri':(9,2,1,0,7,0),'alef':(8,2,1,0,6,0),'lowe':(11,2,0,0,7,0),'khris':(12,2,0,1,6,0),'torasu':(8,2,0,0,5,0),
'luke':(18,1,1,1,4,0),'gort':(16,1,2,1,5,0),'ken':(13,3,1,0,12,0),'mae':(10,3,2,0,11,0),'arthur':(8,3,0,0,8,0),'pelle':(12,3,2,1,10,0),'vankar':(13,3,1,2,8,0),'earnest':(10,3,2,0,9,0),
'hans':(8,2,0,0,11,0),'diane':(10,2,1,0,9,0),'gong':(13,2,2,1,8,0),'amon':(11,3,1,1,11,1),'balbaroy':(10,3,1,0,12,1),'zylo':(14,3,2,0,12,0),'kokichi':(10,3,2,0,11,1),
'domingo':(10,2,1,1,9,1),'guntz':(18,1,2,3,3,0),'musashi':(14,1,3,2,6,0),'hanzou':(10,3,2,0,14,0),'lyle':(11,2,1,0,7,0),'bleu':(12,2,1,1,9,1),'adam':(15,1,1,2,4,0),'jogurt':(4,1,0,0,3,0)}  # PV, Mov(-1: o jogo soma +1), ATQ, DF, AG, Voa

DESC={'lyle':'Cavaleiro que atira à distância (arcos e balistas). Movimento de centauro com alcance de arqueiro, mas frágil e com pouca defesa.','arthur':'Projeto de longo prazo: começa muito fraco, quase sem causar dano, e deve se tornar o melhor cavaleiro.','bleu':'Dragão filhote e voador: entra frágil, mas é o futuro dragão poderoso do grupo.'}
# ---------------- Heróis ----------------
w=wb['Heróis']
w.delete_rows(1,w.max_row)
w.append(['Chave','Herói','Raça','Classe','Descrição','Classe base','PV','Movimento','ATQ','DF','AG','Voa'])
src=openpyxl.load_workbook(A+'9864fbc8-cartas.xlsx',data_only=True)['Heróis']
for r in list(src.iter_rows(values_only=True))[1:]:
    k=r[0]; s=ST[k]
    w.append([k,r[1],r[2],r[3],DESC.get(k,r[4]),BASE[k],s[0],s[1]+1,s[2],s[3],s[4],'Sim' if s[5] else 'Não'])
head(w); widths(w,[11,12,16,22,60,22,6,11,6,6,6,6])
for row in w.iter_rows(min_row=2): row[4].alignment=Alignment(wrap_text=True,vertical='top')
dv=DataValidation(type='list',formula1='"Sim,Não"'); w.add_data_validation(dv); dv.add(f'L2:L{w.max_row}')

# ---------------- Decks Classe / Assinaturas ----------------
del wb['Cartas Heróis']
def actdv(w,cols,n):
    d=DataValidation(type='list',formula1='"Atacar,Mover,Curar,Magia,Saquear"'); w.add_data_validation(d)
    for c in cols: d.add(f'{c}2:{c}{n+1}')
    d2=DataValidation(type='list',formula1='"Sim,Não"'); w.add_data_validation(d2); return d2
w=wb.create_sheet('Decks Classe',2)
w.append(['Classe base','Nº','Nome da carta','Topo: ação','Topo: valor','Base: ação','Base: valor','Queimar','Topo: efeitos','Base: efeitos'])
for cls,cs in DECKS.items():
    for i,s in enumerate(cs): w.append([cls,i+1,*card(s)])
head(w); widths(w,[22,5,24,12,12,12,12,10,28,28])
d2=actdv(w,'DF',w.max_row-1); d2.add(f'H2:H{w.max_row}')
w=wb.create_sheet('Assinaturas',3)
w.append(['Chave','Herói','Nome da carta','Topo: ação','Topo: valor','Base: ação','Base: valor','Queimar','Topo: efeitos','Base: efeitos'])
names={r[0]:r[1] for r in list(src.iter_rows(values_only=True))[1:]}
for k,s in SIG.items(): w.append([k,names[k],*card(s)])
head(w); widths(w,[11,12,24,12,12,12,12,10,28,28])
d2=actdv(w,'DF',w.max_row-1); d2.add(f'H2:H{w.max_row}')

# ---------------- Armas ----------------
CL={
'Short Sword':'Espadachim, Guerreiro, Aviário','Middle Sword':'Espadachim, Guerreiro, Aviário','Long Sword':'Espadachim, Aviário',
'Steel Sword':'Herói, Ninja, Samurai, SKYW','Broad Sword':'Herói, Ninja, Samurai, SKYW','Doom Blade':'','Katana':'',
'Sword of Light':'Herói','Sword of Darkness':'Herói, SKYW','Chaos Breaker':'Herói',
'Hand Axe':'Guerreiro','Middle Axe':'Guerreiro','Battle Axe':'Guerreiro','Heat Axe':'Gladiador','Great Axe':'Gladiador','Atlas':'Gladiador',
'Spear':'Cavaleiro, Cavaleiro a Vapor, Cavaleiro Alado','Bronze Lance':'Cavaleiro, Cavaleiro a Vapor, Cavaleiro Alado','Power Spear':'Cavaleiro, Cavaleiro a Vapor, Cavaleiro Alado',
'Steel Lance':'Paladino, SBRN, SKYL','Chrome Lance':'Paladino, SBRN, SKYL','Halberd':'Paladino, SBRN, SKYL','Devil Lance':'Paladino, SBRN, SKYL','Valkyrie':'Paladino, SBRN, SKYL',
'Wooden Arrow':'Arqueiro, Cavaleiro de Assalto','Steel Arrow':'Arqueiro, Cavaleiro de Assalto','Elven Arrow':'Arqueiro, Cavaleiro de Assalto',
'Assault Shell':'BWMS, SKNT','Buster Shot':'BWMS, SKNT',
'Wooden Staff':'Mago, Sacerdote','Power Staff':'Mago, Sacerdote','Guardian Staff':'Feiticeiro, Vigário','Holy Staff':'Vigário','Demon Rod':'Feiticeiro'}
ONLY={'Doom Blade':'hanzou','Katana':'musashi','Sword of Light':'max','Sword of Darkness':'max','Chaos Breaker':'max'}
PROMO=['Herói', 'Gladiador', 'Paladino', 'SBRN', 'SKYL', 'BWMS', 'SKNT', 'Feiticeiro', 'Vigário', 'SKYW']
CL={k:', '.join(c for c in v.split(', ') if c and c not in PROMO) for k,v in CL.items()}
ONLY={k:v for k,v in ONLY.items() if v!='max'}
w=wb.create_sheet('Armas',4)
w.append(['Nome','Tipo','Alcance','ATQ','Magia','Custo','Classes que podem usar','Personagens exclusivos (chave)','Extras'])
for r in list(arm.iter_rows(values_only=True))[1:]:
    if not r[0]: continue
    n=r[0]; assert n in CL,n
    w.append([n,r[1],str(r[2]),r[3],r[4] if r[4]!='-' else '',r[5],CL[n],ONLY.get(n,''),r[7] or ''])
head(w); widths(w,[20,18,9,6,14,9,48,26,16])

# ---------------- Instruções ----------------
ws=wb['Instruções']; ws.delete_rows(1,ws.max_row)
L=['Banco de dados — Shining Force - Boardgame Versão Digital','',
'Edite as abas e use este arquivo para gerar o index.html (python3 build.py dados.xlsx). Não altere as colunas "Chave" nem os nomes dos monstros.','',
'ABA Heróis: raça, classe (texto exibido), descrição, Classe base (define o deck e as armas permitidas; Sacerdote/Sacerdotisa = "Sacerdote", Cavaleiro/Cavaleira = "Cavaleiro"), PV, Movimento (somado ao Mover das cartas), ATQ (somado ao Atacar das cartas), DF, AG (iniciativa) e Voa.',
'ABA Decks Classe: 7 cartas por Classe base. Heróis da mesma classe base compartilham o mesmo deck.',
'ABA Assinaturas: 1 carta exclusiva por herói, exibida com a cor do herói. Cada herói joga com 8 cartas (7 da classe + 1 de assinatura).',
'Ação: Atacar, Mover, Curar, Magia ou Saquear. Saquear N: recolhe todo o ouro que estiver a até N casas do herói (inclusive na própria casa); o valor da carta é o alcance (hoje, 1). Todo inimigo derrotado deixa 1 de ouro na casa onde morreu, e o herói que termina o turno sobre uma casa com ouro o coleta. Valor: inteiro >= 1 (Mover pode ser 0: o Movimento base do herói é somado). Atacar soma o ATQ do herói + o ATQ da arma equipada. Magia não soma nada (nem ATQ, nem arma): causa o valor da carta, usa o baralho de modificadores e ignora a DF do alvo.',
'Efeitos (colunas "efeitos" de cada metade da carta, separados por ";"): Perfurar N (ignora N de DF), Área (também atinge os inimigos vizinhos do alvo, com o mesmo modificador), Empurrar N (afasta o alvo N casas), Ferir (veneno: 1 de dano por rodada, 2 rodadas), Atordoar (o alvo perde o próximo turno), Congelar (o alvo não pode mover no próximo turno), Escudo N (reduz N de cada dano sofrido até o seu próximo turno; em Curar, vale para quem foi curado), Retaliar N (quem o atacar corpo a corpo sofre N de dano, até o seu próximo turno), Cura em área (Curar atinge todos os aliados feridos a até 2 casas), Salto (Mover atravessa inimigos), Drenar N (cura N PV ao causar dano). Efeitos de ataque não valem se o modificador for Erro.',
'Queimar = Sim: depois de jogada (topo ou base) a carta sai do jogo para sempre. O aviso só aparece na carta quando for Sim.',
'Descansar: ação livre, 1 vez por turno, antes de escolher as cartas ou depois de concluir as ações. O descarte volta para a mão e 1 carta aleatória é queimada. Não cura PV.',
'ABA Armas: Alcance 1 = só o alcance normal (adjacente); 2 = ataca até 2 casas; "2-3" = ataca a 2 e 3 casas, mas não a 1. Todo herói tem alcance 1 sem arma. Classes que podem usar: nomes de Classe base separados por vírgula. Personagens exclusivos: chave do herói (vale mesmo que a classe não conste). Armas de classes de promoção ficam sem classe (ninguém as usa) até o sistema de promoção ser definido.',
'ABA Cartas Monstros: Mover soma o MV do monstro; Atacar soma o AT. Todos do mesmo tipo executam a mesma carta. Embaralhar = Sim marca a carta de embaralhamento (mantenha 1 por baralho).',
'ABA Alcance Monstros: alcance de ataque em casas (1 = só adjacente; a diagonal custa 2) e habilidades (Heal, Sleep, Freeze, Poison...), uma linha por monstro.']
for i,l in enumerate(L): ws.cell(i+1,1,l)
ws['A1'].font=Font(bold=True,size=14); ws.column_dimensions['A'].width=150
for r in range(2,ws.max_row+1): ws.cell(r,1).alignment=Alignment(wrap_text=True,vertical='top')
wb.move_sheet('Heróis',offset=0)
wb.save('dados.xlsx'); print(wb.sheetnames)
