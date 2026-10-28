import re, json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
tpl=open('tpl.html').read()
H=re.findall(r"\['(\w+)', '(\w+)', '(\w+)', '([^']+)', (\d+), (\d+), (\d+), (\d+), (\d+), (\d+), (\d+), (\d+), '([^']*)'\]",tpl)
mons=json.load(open('monsters.json'))
A={'atk':'Atacar','mov':'Mover','heal':'Curar'}
def c1(v): return max(1,v)
def hero_cards(atk,heal):
    a=atk; c=heal
    d=[('Golpe Firme',('atk',a),('mov',2)),('Avanço',('mov',2),('atk',a)),('Investida',('atk',a+1),('mov',1)),('Reposicionar',('mov',3),('atk',a-1))]
    if c: d+=[('Prece',('heal',c),('mov',2)),('Apoio em Combate',('atk',a),('heal',c)),('Bênção',('heal',c+1),('mov',1)),('Luz Restauradora',('heal',c+1),('atk',a-1))]
    else: d+=[('Fúria',('atk',a),('atk',a-1)),('Sprint',('mov',3),('mov',2)),('Ataque Rápido',('atk',a-1),('mov',3)),('Postura Pesada',('atk',a+1),('mov',1))]
    return [(n,(t[0],c1(t[1])),(b[0],c1(b[1])),False) for n,t,b in d]
def mon_cards(sp):
    m=re.search(r'heal (\d)',(sp or '').lower()); h=int(m.group(1))+1 if m else 0
    c=[['Avanço',('mov',2),('atk',2)],['Investida',('mov',3),('atk',1)],['Golpe',('atk',3),('mov',1)],['Pancada',('atk',2),('mov',2)],['Cerco',('mov',1),('atk',3)],['Fúria',('atk',3),('atk',1)],['Passo Curto',('mov',1),('atk',3)]]
    if h: c[3]=['Cura e Golpe',('atk',2),('heal',h)]; c[6]=['Cura',('heal',h),('mov',2)]
    c.append(['Embaralhar',('mov',2),('atk',3)])
    return [(n,t,b,i==7) for i,(n,t,b) in enumerate(c)]
def build(path):
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='Instruções'
    for i,l in enumerate([
     'Cartas de ação — Shining Force - Boardgame Versão Digital',
     '',
     'Edite as abas "Cartas Heróis" e "Cartas Monstros" e devolva o arquivo para atualizar as cartas do jogo.',
     'Ação: use somente Atacar, Mover ou Curar (há lista suspensa). Valor: número inteiro maior ou igual a 1.',
     'Cada carta tem duas ações: Topo e Base. O herói joga o topo de uma carta e a base de outra.',
     'Heróis: os valores são do nível 1. Atacar e Curar ganham +1 por nível acima do 1. Mover soma o movimento base do herói (1, 2 ou 3).',
     'Monstros: Mover soma o MV da planilha de monstros (0 a 3; MV 0 não se move). Atacar soma o AT do monstro.',
     'Todos os monstros do mesmo tipo executam a mesma carta revelada no turno do grupo.',
     'Embaralhar = Sim marca a carta de embaralhamento. Quando ela é revelada, o descarte é reembaralhado no fim da rodada. Mantenha 1 por baralho.',
     'Não altere a coluna "Chave" nem os nomes dos monstros. Pode mudar o nome da carta, as ações, os valores e a quantidade de cartas por baralho.']):
        ws.cell(row=i+1,column=1,value=l)
    ws['A1'].font=Font(bold=True,size=14); ws.column_dimensions['A'].width=130
    hd=PatternFill('solid',fgColor='2D323C'); hf=Font(bold=True,color='FFFFFF')
    def sheet(name,head,rows,widths,actcols):
        w=wb.create_sheet(name); w.append(head)
        for c in range(1,len(head)+1):
            cell=w.cell(row=1,column=c); cell.fill=hd; cell.font=hf; cell.alignment=Alignment(vertical='center')
        for r in rows: w.append(r)
        for i,x in enumerate(widths): w.column_dimensions[get_column_letter(i+1)].width=x
        w.freeze_panes='A2'
        dv=DataValidation(type='list',formula1='"Atacar,Mover,Curar"',allow_blank=False); w.add_data_validation(dv)
        for col in actcols: dv.add(f'{col}2:{col}{len(rows)+1}')
        return w
    hr=[]
    for k,nm,gl,cls,hp,mv,atk,heal,rng,df,fly,desc in H:
        for i,(n,t,b,s) in enumerate(hero_cards(int(atk),int(heal))):
            hr.append([k,nm,i+1,n,A[t[0]],t[1],A[b[0]],b[1]])
    sheet('Cartas Heróis',['Chave','Herói','Nº','Nome da carta','Topo: ação','Topo: valor','Base: ação','Base: valor'],hr,[12,14,5,24,12,12,12,12],['E','G'])
    mr=[]
    for m in mons:
        for i,(n,t,b,s) in enumerate(mon_cards(m['sp'])):
            mr.append([m['n'],i+1,n,A[t[0]],t[1],A[b[0]],b[1],'Sim' if s else 'Não'])
    w=sheet('Cartas Monstros',['Monstro','Nº','Nome da carta','Topo: ação','Topo: valor','Base: ação','Base: valor','Embaralhar'],mr,[28,5,24,12,12,12,12,12],['D','F'])
    dv2=DataValidation(type='list',formula1='"Sim,Não"'); w.add_data_validation(dv2); dv2.add(f'H2:H{len(mr)+1}')
    wb.save(path)
if __name__=='__main__':
    build('cartas.xlsx'); print('heroes',len(H),'monsters',len(mons))
